import os
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from datetime import datetime

# ================== 【关键：根据截图精准配置】==================
ORIGINAL_FILE = r"D:\project\py\mnt\checklist-统一业务监控系统V26.01.01版本变更实施.xlsx"
SHEET_MAIN = "checklist模板"  # 主sheet名称
SHEET_AOPS = "AOPS变更编排"  # 第二个AOPS sheet名称

# 🔑 请务必按截图确认（已根据您最新截图校准）：
CELL_CHANGE_TITLE = "B5"  # 变更标题的值（截图中显示"统一业务监控系统V26.01.02"的位置）
CELL_START_TIME = "H5"  # 开始实施时间的值（截图中"2026-01-29 22:30"所在单元格）
CELL_WORK_ORDER = "C6"  # 变更工单号的值（截图中工单号所在单元格）
CELL_REF_TIME = "F20"  # 参考用时单元格
CELL_SQL_STATEMENT = "C29"  # SQL语句单元格
CELL_AOPS_DESC = "C32"  # AOPS发布单说明单元格
CELL_AOPS_TIME = "B4"  # AOPS sheet实施时间单元格
CELL_AOPS_ORDER = "C4"  # AOPS sheet发布单单元格

# 传入的新值（新增发布单字段，与原有字段同区域）
NEW_CHANGE_TITLE = "统一业务监控系统V26.01.02"
NEW_START_TIME = "开始实施时间：2026-01-30 22:30"
NEW_WORK_ORDER = "CM20260130105841291914"
NEW_REF_TIME = 20  # 参考用时（数字，自动拼接"分钟"）
NEW_AOPS_ORDER = "R2026012913430962719811"  # 发布单字段（新增）


# ==========================================================

def safe_set_cell_value(sheet, cell_addr, value):
    """
    安全设置单元格值（自动处理合并单元格）
    :param sheet: 工作表对象
    :param cell_addr: 单元格地址字符串，如 "G5"
    :param value: 要设置的值
    """
    # 检查是否在合并区域内
    for merged_range in sheet.merged_cells.ranges:
        if cell_addr in merged_range:
            # 修改合并区域左上角单元格
            start_row = merged_range.min_row
            start_col = merged_range.min_col
            sheet.cell(row=start_row, column=start_col).value = value
            print(f"✓ 已修改合并区域 [{merged_range}] 起始单元格 {get_column_letter(start_col)}{start_row} = {value}")
            return
    # 非合并单元格：直接赋值
    sheet[cell_addr] = value
    print(f"✓ 已直接修改单元格 {cell_addr} = {value}")


def generate_new_filename(old_path, new_ver):
    """用新版本号替换原文件名中的旧版本号"""
    base = os.path.basename(old_path)
    # 严格替换第一个匹配的版本号
    new_base = re.sub(r'V\d+\.\d+\.\d+', new_ver, base, count=1)
    return os.path.join(os.path.dirname(old_path), new_base)


def parse_start_time_id(start_time_str):
    """
    从NEW_START_TIME解析出20260129_1格式的id
    入参：start_time_str（如"开始实施时间：2026-01-30 22:30"）
    出参：拼接后的id（如"20260130_1"）
    """
    # 正则匹配日期部分（yyyy-mm-dd）
    date_match = re.search(r'(\d{4})-(\d{2})-(\d{2})', start_time_str)
    if not date_match:
        raise ValueError(f"从开始实施时间【{start_time_str}】解析日期失败，格式要求包含 yyyy-mm-dd")
    year, month, day = date_match.groups()
    return f"{year}{month}{day}_1"  # 拼接为yyyyMMdd_1格式


def parse_month_desc(change_title):
    """
    从NEW_CHANGE_TITLE解析出"X月迭代版本"描述
    入参：change_title（如"统一业务监控系统V26.01.02"）
    出参：月份描述（如"1月迭代版本"）
    """
    # 正则匹配版本号中的月份（Vxx.MM.xx 中的MM）
    month_match = re.search(r'V\d+\.(\d{2})\.', change_title)
    if not month_match:
        raise ValueError(f"从变更标题【{change_title}】解析月份失败，版本号格式要求 Vxx.MM.xx")
    month = int(month_match.group(1))  # 去除前导零（如01→1）
    return f"{month}月迭代版本"


def generate_sql_statement():
    """根据各参数动态生成C29需要的SQL语句"""
    # 解析各动态参数
    sql_id = parse_start_time_id(NEW_START_TIME)  # 20260130_1
    sql_viminal_no = NEW_WORK_ORDER  # 变更工单号
    sql_viminal_name = NEW_CHANGE_TITLE  # 变更标题
    # 从变更标题解析版本号（与生成新文件名的解析逻辑一致）
    ver_match = re.search(r'V\d+\.\d+\.\d+', NEW_CHANGE_TITLE)
    if not ver_match:
        raise ValueError("变更标题中未找到有效版本号（格式如 V26.01.02）")
    sql_db_version = ver_match.group()  # V26.01.02
    sql_descs = parse_month_desc(NEW_CHANGE_TITLE)  # 1月迭代版本
    # 固定值（可根据需要修改）
    sql_developer = "李立伟"

    # 拼接SQL语句（保持原格式，动态替换变量）
    sql = (
        f"INSERT INTO sleyedb.log_db_version (id, viminal_no, viminal_name, app_version, db_version, descs, developer, create_time) "
        f"VALUES ('{sql_id}', '{sql_viminal_no}', '{sql_viminal_name}', '', '{sql_db_version}', '{sql_descs}', '{sql_developer}', NOW());"
    )
    return sql


def parse_aops_time(start_time_str):
    """
    从NEW_START_TIME解析并转换为AOPS sheet需要的时间格式
    入参：start_time_str（如"开始实施时间：2026-01-30 22:30"）
    出参：格式化时间（如"2026/1/30 22:30:00"）
    """
    # 正则匹配完整时间（yyyy-mm-dd hh:mm）
    time_match = re.search(r'(\d{4})-(\d{2})-(\d{2})\s+(\d{2}):(\d{2})', start_time_str)
    if not time_match:
        raise ValueError(f"从开始实施时间【{start_time_str}】解析时间失败，格式要求包含 yyyy-mm-dd hh:mm")
    year, month, day, hour, minute = time_match.groups()
    # 转换为指定格式（去除月份/日期前导零，拼接秒为00）
    return f"{year}/{int(month)}/{int(day)} {hour}:{minute}:00"


try:
    # 1. 基础解析：新版本号（用于生成新文件名）
    ver_match = re.search(r'V\d+\.\d+\.\d+', NEW_CHANGE_TITLE)
    if not ver_match:
        raise ValueError("变更标题中未找到有效版本号（格式如 V26.01.02）")
    new_version = ver_match.group()
    NEW_FILE = generate_new_filename(ORIGINAL_FILE, new_version)

    # 2. 加载工作簿（关键：必须 read_only=False，支持写入）
    wb = load_workbook(ORIGINAL_FILE, read_only=False, data_only=False)

    # ================== 处理主sheet：checklist模板 ==================
    if SHEET_MAIN not in wb.sheetnames:
        raise ValueError(f"主工作表 '{SHEET_MAIN}' 不存在！当前工作表: {wb.sheetnames}")
    sheet_main = wb[SHEET_MAIN]

    # 3. 原有字段更新（保持原逻辑）
    safe_set_cell_value(sheet_main, CELL_CHANGE_TITLE, NEW_CHANGE_TITLE)
    safe_set_cell_value(sheet_main, CELL_START_TIME, NEW_START_TIME)
    safe_set_cell_value(sheet_main, CELL_WORK_ORDER, NEW_WORK_ORDER)

    # 4. 新增字段1：参考用时（自动拼接"分钟"）
    ref_time_value = f"{NEW_REF_TIME}分钟"
    safe_set_cell_value(sheet_main, CELL_REF_TIME, ref_time_value)

    # 5. 新增字段2：C29 SQL语句（动态生成）
    sql_value = generate_sql_statement()
    safe_set_cell_value(sheet_main, CELL_SQL_STATEMENT, sql_value)

    # 6. 新增字段3：C32 AOPS发布单说明（固定格式+动态发布单）
    aops_desc_value = (
        f"在AOPS中执行发布单为{NEW_AOPS_ORDER}的变更，流水线类型：详见标签页AOPS变更编排\n"
        "结果信息查看：（变更计划实际执行结果为部署管理-部署流水线中实际执行）\n"
        "1、在变更计划执行列表中，操作列，点击跳转详情，查看流水线实际执行情况；\n"
        "2、部署管理-部署流水线列表，查找对应的流水线运行情况，打开节点查看详细日志"
    )
    safe_set_cell_value(sheet_main, CELL_AOPS_DESC, aops_desc_value)

    # ================== 处理第二个sheet：AOPS变更编排 ==================
    if SHEET_AOPS not in wb.sheetnames:
        raise ValueError(f"AOPS工作表 '{SHEET_AOPS}' 不存在！当前工作表: {wb.sheetnames}")
    sheet_aops = wb[SHEET_AOPS]

    # 7. AOPS sheet B4：实施时间（解析+格式转换）
    aops_time_value = parse_aops_time(NEW_START_TIME)
    safe_set_cell_value(sheet_aops, CELL_AOPS_TIME, aops_time_value)

    # 8. AOPS sheet C4：发布单（直接取NEW_AOPS_ORDER参数）
    safe_set_cell_value(sheet_aops, CELL_AOPS_ORDER, NEW_AOPS_ORDER)

    # 9. 保存为新文件并关闭工作簿
    wb.save(NEW_FILE)
    wb.close()

    # 打印成功信息
    print("\n✅ 操作成功！")
    print(f"📄 原文件: {ORIGINAL_FILE}")
    print(f"💾 新文件: {NEW_FILE}")
    print(f"✨ 已更新所有字段:")
    print(f"   - 变更标题 ({CELL_CHANGE_TITLE}): {NEW_CHANGE_TITLE}")
    print(f"   - 开始实施时间 ({CELL_START_TIME}): {NEW_START_TIME}")
    print(f"   - 变更工单号 ({CELL_WORK_ORDER}): {NEW_WORK_ORDER}")
    print(f"   - 参考用时 ({CELL_REF_TIME}): {ref_time_value}")
    print(f"   - SQL语句 ({CELL_SQL_STATEMENT}): {sql_value}")
    print(f"   - AOPS说明 ({CELL_AOPS_DESC}): 已生成对应发布单说明")
    print(f"   - {SHEET_AOPS}实施时间 ({CELL_AOPS_TIME}): {aops_time_value}")
    print(f"   - {SHEET_AOPS}发布单 ({CELL_AOPS_ORDER}): {NEW_AOPS_ORDER}")

except Exception as e:
    print(f"\n❌ 错误: {type(e).__name__} - {str(e)}")
    import traceback

    traceback.print_exc()